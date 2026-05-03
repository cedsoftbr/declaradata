from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError


def get_client_ip(request) -> str | None:
    xff = request.META.get('HTTP_X_FORWARDED_FOR', '')
    if xff:
        return xff.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR') or None


def read_excel_file(file_obj: Any) -> list[dict[str, Any]]:
    """
    Lê um arquivo Excel e retorna uma lista de dicionários baseada no cabeçalho.
    """
    try:
        workbook = load_workbook(filename=BytesIO(file_obj.read()), data_only=True)
        sheet = workbook.active

        if not sheet:
            raise ValidationError('Arquivo Excel vazio ou sem planilhas ativas.')

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        header = [str(cell).strip().lower() for cell in rows[0]]
        data = []

        for row in rows[1:]:
            if any(cell is not None for cell in row):
                row_dict = dict(zip(header, row, strict=False))
                data.append(row_dict)

        return data
    except Exception as e:
        raise ValidationError(f'Erro ao ler arquivo Excel: {e}') from e
