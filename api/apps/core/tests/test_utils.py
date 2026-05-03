from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError

from apps.core.utils import read_excel_file


def test_read_excel_file_success():
    wb = Workbook()
    ws = wb.active
    ws.append(['Nome', 'IBGE'])
    ws.append(['Test', '1234567'])

    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    file_obj = SimpleUploadedFile('test.xlsx', file_stream.read())
    data = read_excel_file(file_obj)

    assert len(data) == 1
    assert data[0]['nome'] == 'Test'
    assert str(data[0]['ibge']) == '1234567'


def test_read_excel_file_no_data_only_header():
    wb = Workbook()
    ws = wb.active
    ws.append(['Nome', 'IBGE'])
    # Only header, max_row will be 1
    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    file_obj = SimpleUploadedFile('test.xlsx', file_stream.read())
    data = read_excel_file(file_obj)
    assert data == []


def test_read_excel_file_with_empty_row():
    wb = Workbook()
    ws = wb.active
    ws.append(['Nome', 'IBGE'])
    ws.append(['Test', '1234567'])
    ws.append([None, None])  # Empty row
    ws.append(['Test 2', '2222222'])

    file_stream = BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)
    file_obj = SimpleUploadedFile('test.xlsx', file_stream.read())
    data = read_excel_file(file_obj)
    expected_len = 2
    assert len(data) == expected_len
    assert data[1]['nome'] == 'Test 2'


def test_read_excel_file_invalid():
    file_obj = SimpleUploadedFile('test.txt', b'not an excel file')
    with pytest.raises(ValidationError, match='Erro ao ler arquivo Excel'):
        read_excel_file(file_obj)
